# Copyright 2025 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Unit tests for workbook routing and Spanner Graph planning. No live services."""

import io

import pandas as pd
import pytest
from openpyxl import Workbook

from data_science.sub_agents.database import xlsx_loader


def _stub_tables(monkeypatch):
    monkeypatch.setattr(
        xlsx_loader,
        "load_table",
        lambda workbook_name, sheet, frame: {
            "sheet": sheet,
            "table": f"p.d.{sheet}",
            "table_id": sheet,
            "row_count": len(frame),
        },
    )
    monkeypatch.setattr(xlsx_loader, "drop_bq_tables", lambda tables: None)
    monkeypatch.setattr(xlsx_loader, "_previous_tables", lambda name: [])
    monkeypatch.setattr(xlsx_loader, "_drop_previous_graph", lambda name: None)


def _workbook() -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(
            {"person_id": [1, 2], "name": ["Ada", "Gus"]}
        ).to_excel(writer, sheet_name="People", index=False)
        pd.DataFrame(
            {"source": [1], "target": [2], "since": ["2020"]}
        ).to_excel(writer, sheet_name="Knows", index=False)
        pd.DataFrame({"note": ["ignore me"]}).to_excel(
            writer, sheet_name="Notes", index=False
        )
    return buffer.getvalue()


def test_declared_mapping_builds_one_graph_and_ignores_undeclared_sheets(monkeypatch):
    _stub_tables(monkeypatch)
    captured = {}

    class FakeDb:
        def update_ddl(self, statements):
            captured["ddl"] = statements

            class Op:
                def result(self, timeout=0):
                    return None

            return Op()

        def run_in_transaction(self, fn):
            captured.setdefault("writes", []).append(fn)

        def snapshot(self):
            raise AssertionError("schema lookup is not part of this test")

    monkeypatch.setattr(xlsx_loader, "database", lambda: FakeDb())
    monkeypatch.setattr(xlsx_loader, "ensure_catalog", lambda db: None)
    monkeypatch.setattr(xlsx_loader, "insert_rows", lambda *args: None)
    monkeypatch.setattr(xlsx_loader, "graph_schema", lambda name: None)

    result = xlsx_loader.load_workbook(
        _workbook(),
        "friends.xlsx",
        node_sheets=[{"sheet": "People", "id_column": "person_id"}],
        edge_sheet={
            "sheet": "Knows",
            "source_column": "source",
            "target_column": "target",
            "source_node": "People",
            "target_node": "People",
        },
    )

    assert result["status"] == "SUCCESS"
    assert result["graph_name"].startswith("g_friends_")
    ddl = "\n".join(captured["ddl"])
    assert "CREATE OR REPLACE PROPERTY GRAPH" in ddl
    assert "SOURCE KEY (source) REFERENCES" in ddl
    assert "DESTINATION KEY (target) REFERENCES" in ddl
    assert "Notes" not in ddl
    assert "CREATE TABLE IF NOT EXISTS workbook_catalog" not in ddl


def test_missing_edge_declaration_does_not_infer_edges(monkeypatch):
    _stub_tables(monkeypatch)
    class FakeDb:
        def update_ddl(self, statements):
            self.statements = statements

            class Op:
                def result(self, timeout=0):
                    return None

            return Op()

        def run_in_transaction(self, fn):
            return None

    fake = FakeDb()
    monkeypatch.setattr(xlsx_loader, "database", lambda: fake)
    monkeypatch.setattr(xlsx_loader, "ensure_catalog", lambda db: None)
    monkeypatch.setattr(xlsx_loader, "insert_rows", lambda *args: None)
    monkeypatch.setattr(xlsx_loader, "graph_schema", lambda name: None)
    xlsx_loader.load_workbook(
        _workbook(),
        "friends.xlsx",
        node_sheets=[{"sheet": "People", "id_column": "person_id"}],
    )
    ddl = "\n".join(fake.statements)
    assert "EDGE TABLES" not in ddl
    assert "Knows" not in ddl


def test_inference_loads_nodes_and_self_edge(monkeypatch):
    _stub_tables(monkeypatch)
    captured = {}

    class FakeDb:
        def update_ddl(self, statements):
            captured["ddl"] = statements

            class Op:
                def result(self, timeout=0):
                    return None

            return Op()

        def run_in_transaction(self, fn):
            return None

    monkeypatch.setattr(xlsx_loader, "database", lambda: FakeDb())
    monkeypatch.setattr(xlsx_loader, "ensure_catalog", lambda db: None)
    monkeypatch.setattr(xlsx_loader, "insert_rows", lambda *args: None)
    monkeypatch.setattr(xlsx_loader, "graph_schema", lambda name: None)

    result = xlsx_loader.load_workbook(_workbook(), "friends.xlsx")

    assert result["inferred"] is True
    assert result["confidence"] == "medium"
    nodes = {node["sheet"]: node for node in result["mapping"]["nodes"]}
    assert nodes["People"]["id_column"] == "person_id"
    assert nodes["People"]["confidence"] == "high"
    assert "Notes" not in nodes
    assert result["tables"][0]["sheet"] == "Notes"
    edges = result["mapping"]["edges"]
    assert len(edges) == 1
    assert edges[0]["sheet"] == "Knows"
    assert edges[0]["source_node"] == "People"
    assert edges[0]["target_node"] == "People"
    assert edges[0]["confidence"] == "high"
    ddl = "\n".join(captured["ddl"])
    assert "EDGE TABLES" in ddl
    assert "Knows" in ddl
    assert "Notes" not in ddl


def test_second_edge_and_name_match_beat_overlap():
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"id": [1, 2, 3, 4, 5]}).to_excel(
            writer, sheet_name="Person", index=False
        )
        pd.DataFrame({"id": [1, 2, 3, 4]}).to_excel(
            writer, sheet_name="Group", index=False
        )
        pd.DataFrame(
            {"person_id": [1, 2, 3, 4, 5], "group_id": [1, 2, 3, 4, 1]}
        ).to_excel(writer, sheet_name="Member", index=False)
        pd.DataFrame({"left_id": [1, 2], "right_id": [2, 3]}).to_excel(
            writer, sheet_name="Follows", index=False
        )
    analysis = xlsx_loader.analyze_workbook(buffer.getvalue())
    edges = {edge["sheet"]: edge for edge in analysis["edge_sheets"]}
    assert edges["Member"]["source_node"] == "Person"
    assert edges["Member"]["target_node"] == "Group"
    assert edges["Follows"]["source_node"] == "Person"
    assert {spec["sheet"] for spec in analysis["node_sheets"]} == {"Person", "Group"}


def test_explicit_mapping_overrides_inference(monkeypatch):
    _stub_tables(monkeypatch)
    class FakeDb:
        def update_ddl(self, statements):
            self.statements = statements

            class Op:
                def result(self, timeout=0):
                    return None

            return Op()

        def run_in_transaction(self, fn):
            return None

    fake = FakeDb()
    monkeypatch.setattr(xlsx_loader, "database", lambda: fake)
    monkeypatch.setattr(xlsx_loader, "ensure_catalog", lambda db: None)
    monkeypatch.setattr(xlsx_loader, "insert_rows", lambda *args: None)
    monkeypatch.setattr(xlsx_loader, "graph_schema", lambda name: None)

    result = xlsx_loader.load_workbook(
        _workbook(),
        "friends.xlsx",
        node_sheets=[{"sheet": "People", "id_column": "person_id"}],
    )

    assert result["inferred"] is False
    assert result["confidence"] == "declared"
    assert result["mapping"]["edges"] == []
    ddl = "\n".join(fake.statements)
    assert "EDGE TABLES" not in ddl
    assert "Knows" not in ddl


def test_low_confidence_still_loads_with_synthetic_key():
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"city": ["Paris", "Paris"], "note": ["a", "a"]}).to_excel(
            writer, sheet_name="Places", index=False
        )
    analysis = xlsx_loader.analyze_workbook(buffer.getvalue())
    assert analysis["confidence"] == "low"
    assert analysis["node_sheets"][0]["id_column"] is None
    sheets = xlsx_loader.read_sheets(buffer.getvalue())
    plan = xlsx_loader._node_plan(
        sheets, {"sheet": "Places", "id_column": None}, "g_places"
    )
    assert plan["synthetic_id"] is True
    assert plan["frame"]["row_id"].is_unique


def test_insert_rows_commits_edge_ids():
    class FakeTxn:
        def __init__(self):
            self.calls = []

        def insert_or_update(self, table, columns, values):
            self.calls.append(("upsert", table, columns, values))

        def insert(self, table, columns, values):
            self.calls.append(("insert", table, columns, values))

    class FakeDb:
        def __init__(self):
            self.txn = FakeTxn()

        def run_in_transaction(self, fn):
            fn(self.txn)

    people = pd.DataFrame({"person_id": [1], "name": ["Ada"]})
    knows = pd.DataFrame({"source": [1], "target": [1]})
    node = xlsx_loader._node_plan(
        {"People": people},
        {"sheet": "People", "id_column": "person_id"},
        "g_test",
    )
    edge = xlsx_loader._edge_plan(
        {"Knows": knows},
        {
            "sheet": "Knows",
            "source_column": "source",
            "target_column": "target",
            "source_node": "People",
            "target_node": "People",
        },
        [node],
        "g_test",
    )
    db = FakeDb()
    xlsx_loader.insert_rows(db, [node], [edge])
    assert [call[0] for call in db.txn.calls] == ["upsert", "insert"]
    assert db.txn.calls[1][2][0] == "edge_id"
    assert db.txn.calls[1][3][0][0]


def test_reload_drops_tables_omitted_by_the_new_plan(monkeypatch):
    monkeypatch.setattr(
        xlsx_loader,
        "graph_schema",
        lambda name: {
            "mapping": {
                "nodes": [{"table": "g_friends_People"}],
                "edges": [{"table": "g_friends_Knows"}],
            }
        },
    )
    drops = xlsx_loader.drop_stale_tables(
        "g_friends",
        [{"table": "g_friends_People"}],
        [],
    )
    assert drops == ["DROP TABLE IF EXISTS g_friends_Knows"]


def test_missing_id_column_is_rejected():
    with pytest.raises(ValueError, match="id column"):
        xlsx_loader._node_plan(
            {"People": pd.DataFrame({"name": ["Ada"]})},
            {"sheet": "People", "id_column": "person_id"},
            "g_test",
        )


def test_guard_sql_allows_graph_and_blocks_dml():
    assert xlsx_loader.guard_sql(
        "GRAPH g MATCH (n) RETURN n LIMIT 5"
    ).startswith("GRAPH")
    with pytest.raises(ValueError):
        xlsx_loader.guard_sql("DROP PROPERTY GRAPH g")
        xlsx_loader.guard_sql("SELECT 1; SELECT 2")


    with pytest.raises(ValueError):
        xlsx_loader.guard_sql("SELECT 1; SELECT 2")


def _formula_book() -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook()
    people = workbook.active
    people.title = "People"
    people.append(["person_id", "name"])
    people.append([1, "Ada"])
    people.append([2, "Grace"])
    notes = workbook.create_sheet("Notes")
    notes.append(["note", "score"])
    notes.append(["a", "=People!A2"])
    notes.append(["b", "=People!A3"])
    raw = workbook.create_sheet("Events")
    raw.append(["event_id", "city"])
    for index in range(3):
        raw.append([index, "Paris"])
    workbook.save(buffer)
    return buffer.getvalue()


def test_formula_sheet_routes_to_graph_and_raw_sheet_to_table():
    classified = {
        item["sheet"]: item
        for item in xlsx_loader.classify_sheets(_formula_book())["sheets"]
    }
    assert classified["Notes"]["destination"] == "graph"
    assert classified["Notes"]["formula_count"] == 2
    assert classified["People"]["destination"] == "graph"
    assert classified["Events"]["destination"] == "table"


def test_large_raw_sheet_is_not_rejected(monkeypatch):
    frame = pd.DataFrame({"id": range(xlsx_loader.GRAPH_ROW_HINT + 1), "city": "x"})
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Events", index=False)
    payload = buffer.getvalue()
    classified = xlsx_loader.classify_sheets(payload)["sheets"]
    assert classified[0]["destination"] == "table"
    assert classified[0]["row_count"] == xlsx_loader.GRAPH_ROW_HINT + 1

    captured = {}

    def fake_load(workbook_name, sheet, loaded):
        captured["rows"] = len(loaded)
        return {
            "sheet": sheet,
            "table": f"p.d.{sheet}",
            "table_id": sheet,
            "row_count": len(loaded),
        }

    monkeypatch.setattr(xlsx_loader, "load_table", fake_load)
    monkeypatch.setattr(xlsx_loader, "_previous_tables", lambda name: [])
    monkeypatch.setattr(xlsx_loader, "drop_bq_tables", lambda tables: None)
    monkeypatch.setattr(xlsx_loader, "_drop_previous_graph", lambda name: None)
    result = xlsx_loader.load_workbook(payload, "events.xlsx")
    assert captured["rows"] == xlsx_loader.GRAPH_ROW_HINT + 1


class _FakeDb:
    def update_ddl(self, statements):
        class Op:
            def result(self, timeout=0):
                return None

        return Op()

    def run_in_transaction(self, fn):
        return None


def _patch_graph(monkeypatch, previous_tables=None):
    monkeypatch.setattr(xlsx_loader, "database", lambda: _FakeDb())
    monkeypatch.setattr(xlsx_loader, "ensure_catalog", lambda db: None)
    monkeypatch.setattr(xlsx_loader, "insert_rows", lambda *args: None)
    monkeypatch.setattr(xlsx_loader, "drop_stale_tables", lambda *args: [])
    monkeypatch.setattr(xlsx_loader, "_previous_tables", lambda name: previous_tables or [])


def test_mixed_workbook_splits_destinations(monkeypatch):
    loaded = {}

    def fake_load(workbook_name, sheet, frame):
        loaded[sheet] = len(frame)
        return {
            "sheet": sheet,
            "table": f"p.d.friends_{sheet}",
            "table_id": f"friends_{sheet}",
            "row_count": len(frame),
        }

    monkeypatch.setattr(xlsx_loader, "load_table", fake_load)
    dropped = {}
    monkeypatch.setattr(
        xlsx_loader,
        "drop_bq_tables",
        lambda tables: dropped.setdefault("tables", tables),
    )
    monkeypatch.setattr(
        xlsx_loader,
        "graph_schema",
        lambda name: {"mapping": {"tables": ["p.d.friends_Old"]}},
    )
    _patch_graph(monkeypatch, ["p.d.friends_Old"])

    buffer = io.BytesIO()
    workbook = Workbook()
    people = workbook.active
    people.title = "People"
    people.append(["person_id", "name"])
    people.append([1, "Ada"])
    people.append([2, "Grace"])
    knows = workbook.create_sheet("Knows")
    knows.append(["source", "target"])
    knows.append([1, 2])
    lookup = workbook.create_sheet("Cities")
    lookup.append(["city"])
    lookup.append(["Paris"])
    workbook.save(buffer)

    result = xlsx_loader.load_workbook(buffer.getvalue(), "friends.xlsx")
    assert result["storage"] == "split"
    assert loaded == {"Cities": 1}
    assert result["node_tables"]
    assert dropped["tables"] == ["p.d.friends_Old"]


def test_explicit_mapping_forces_flat_sheet_into_graph(monkeypatch):
    calls = {}

    def fake_load(workbook_name, sheet, frame):
        calls.setdefault("tables", []).append(sheet)
        return {
            "sheet": sheet,
            "table": f"p.d.book_{sheet}",
            "table_id": f"book_{sheet}",
            "row_count": len(frame),
        }

    monkeypatch.setattr(xlsx_loader, "load_table", fake_load)
    monkeypatch.setattr(xlsx_loader, "graph_schema", lambda name: None)
    monkeypatch.setattr(xlsx_loader, "drop_bq_tables", lambda tables: None)
    _patch_graph(monkeypatch)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"city": ["Paris"]}).to_excel(writer, sheet_name="Cities", index=False)
        pd.DataFrame({"note": ["a"]}).to_excel(writer, sheet_name="Notes", index=False)
    result = xlsx_loader.load_workbook(
        buffer.getvalue(),
        "book.xlsx",
        node_sheets=[{"sheet": "Cities"}],
    )
    assert calls["tables"] == ["Notes"]
    assert result["storage"] == "split"
    assert result["inferred"] is False

